#!/usr/bin/env python3
"""
Convert markdown file to Word document (.docx)
"""
import sys
import re
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

def create_document_from_markdown(md_file_path, output_path):
    """Convert markdown to Word document"""

    # Read markdown file
    with open(md_file_path, 'r', encoding='utf-8') as f:
        content = f.read()

    # Create document
    doc = Document()

    # Process line by line
    lines = content.split('\n')
    i = 0
    in_code_block = False
    code_content = []
    in_table = False
    table_rows = []

    while i < len(lines):
        line = lines[i]

        # Handle code blocks
        if line.startswith('```'):
            if not in_code_block:
                in_code_block = True
                code_content = []
            else:
                in_code_block = False
                if code_content:
                    p = doc.add_paragraph()
                    run = p.add_run('\n'.join(code_content))
                    run.font.name = 'Courier New'
                    run.font.size = Pt(9)
                    p.paragraph_format.left_indent = Inches(0.5)
                code_content = []
            i += 1
            continue

        if in_code_block:
            code_content.append(line)
            i += 1
            continue

        # Handle headings
        if line.startswith('# '):
            doc.add_heading(line[2:], level=1)
        elif line.startswith('## '):
            doc.add_heading(line[3:], level=2)
        elif line.startswith('### '):
            doc.add_heading(line[4:], level=3)
        elif line.startswith('#### '):
            doc.add_heading(line[5:], level=4)

        # Handle horizontal rules
        elif line.strip() == '---':
            doc.add_paragraph('_' * 60)

        # Handle tables
        elif '|' in line and line.strip().startswith('|'):
            if not in_table:
                in_table = True
                table_rows = []
            table_rows.append(line)

            # Check if next line is separator or end of table
            if i + 1 < len(lines):
                next_line = lines[i + 1]
                if not ('|' in next_line and next_line.strip().startswith('|')):
                    process_table(doc, table_rows)
                    in_table = False
                    table_rows = []
            else:
                process_table(doc, table_rows)
                in_table = False

        # Handle bullet points
        elif line.strip().startswith('- ') or line.strip().startswith('* '):
            text = line.strip()[2:]
            doc.add_paragraph(text, style='List Bullet')

        # Handle numbered lists
        elif re.match(r'^\d+\.\s', line.strip()):
            text = re.sub(r'^\d+\.\s', '', line.strip())
            doc.add_paragraph(text, style='List Number')

        # Handle checkboxes
        elif line.strip().startswith('- [ ]') or line.strip().startswith('- [x]'):
            checked = '[x]' in line
            text = line.strip()[5:].strip()
            p = doc.add_paragraph()
            p.add_run('☐ ' if not checked else '☑ ').bold = True
            p.add_run(text)

        # Regular paragraphs
        elif line.strip() and not line.startswith('#'):
            if not (line.strip().startswith('|') and set(line.replace('|', '').strip()) <= {'-', ' ', ':'}):
                add_formatted_paragraph(doc, line)

        # Empty line
        else:
            if line.strip() == '':
                doc.add_paragraph()

        i += 1

    # Save document
    doc.save(output_path)
    print(f"Document saved to: {output_path}")

def process_table(doc, table_rows):
    """Process markdown table and add to document"""
    if len(table_rows) < 2:
        return

    rows = []
    for row in table_rows:
        cells = [cell.strip() for cell in row.split('|')[1:-1]]
        if all(set(cell.strip()) <= {'-', ':', ' '} for cell in cells):
            continue
        rows.append(cells)

    if not rows:
        return

    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = 'Light Grid Accent 1'

    for i, row_data in enumerate(rows):
        for j, cell_data in enumerate(row_data):
            cell = table.rows[i].cells[j]
            cell.text = cell_data
            if i == 0:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.bold = True

def add_formatted_paragraph(doc, text):
    """Add paragraph with markdown formatting"""
    p = doc.add_paragraph()
    parts = []
    current = ""
    i = 0

    while i < len(text):
        if text[i:i+2] == '**':
            if current:
                parts.append(('normal', current))
                current = ""
            end = text.find('**', i+2)
            if end != -1:
                parts.append(('bold', text[i+2:end]))
                i = end + 2
                continue
        elif text[i] == '`':
            if current:
                parts.append(('normal', current))
                current = ""
            end = text.find('`', i+1)
            if end != -1:
                parts.append(('code', text[i+1:end]))
                i = end + 1
                continue

        current += text[i]
        i += 1

    if current:
        parts.append(('normal', current))

    for fmt, content in parts:
        run = p.add_run(content)
        if fmt == 'bold':
            run.bold = True
        elif fmt == 'code':
            run.font.name = 'Courier New'
            run.font.size = Pt(10)

if __name__ == '__main__':
    input_file = r'C:\Users\jameelaesa\ACS-Transition-Agent-v0\migration-guides\email\email-service-migration.md'
    output_file = r'C:\Users\jameelaesa\ACS-Transition-Agent-v0\email-service-migration.docx'

    try:
        create_document_from_markdown(input_file, output_file)
        print("Conversion successful!")
    except Exception as e:
        print(f"Error: {e}")
        import traceback
        traceback.print_exc()
        sys.exit(1)
