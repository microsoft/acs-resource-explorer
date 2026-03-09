from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

prs = Presentation()
prs.slide_width = Inches(13.33)
prs.slide_height = Inches(7.5)

BLUE       = RGBColor(0, 120, 212)
DARK       = RGBColor(32, 31, 30)
WHITE      = RGBColor(255, 255, 255)
LIGHT_GRAY = RGBColor(243, 242, 241)
RED        = RGBColor(196, 49, 75)
ORANGE     = RGBColor(209, 115, 0)
GRAY_TEXT  = RGBColor(120, 120, 120)
LIGHT_TEXT = RGBColor(180, 180, 180)
LINK_BLUE  = RGBColor(0, 180, 255)

def blank_slide(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])

def bg(slide, color):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color

def txt(slide, text, l, t, w, h, size=14, bold=False, color=DARK, align=PP_ALIGN.LEFT):
    tb = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color

def rect(slide, l, t, w, h, color):
    s = slide.shapes.add_shape(1, Inches(l), Inches(t), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = color
    s.line.fill.background()
    return s

# ── Slide 1: Title ───────────────────────────────────────────────────────────
s = blank_slide(prs)
bg(s, DARK)
rect(s, 0, 0, 0.35, 7.5, BLUE)
txt(s, "ACS Transition Agent", 0.6, 1.8, 12, 1.2, size=44, bold=True, color=WHITE)
txt(s, "Automated Azure Communication Services Deprecation Assessment", 0.6, 3.1, 11, 0.7, size=22, color=LIGHT_TEXT)
txt(s, "github.com/jameelaesa/ACS-Transition-Agent-v0", 0.6, 6.5, 10, 0.5, size=14, color=GRAY_TEXT)

# ── Slide 2: The Problem ─────────────────────────────────────────────────────
s = blank_slide(prs)
bg(s, WHITE)
rect(s, 0, 0, 13.33, 1.1, BLUE)
txt(s, "The Problem", 0.5, 0.15, 12, 0.8, size=32, bold=True, color=WHITE)
rect(s, 0.5, 1.4, 5.8, 4.8, LIGHT_GRAY)
txt(s, "Microsoft is retiring Azure Communication Services", 0.7, 1.55, 5.4, 0.7, size=18, bold=True, color=DARK)
txt(s, "5 channels are being retired or changed effective March 31, 2029.\n\nOrganizations using ACS must identify which services they use and migrate before the deadline or face service disruption.", 0.7, 2.3, 5.3, 2.8, size=13, color=DARK)
rect(s, 6.8, 1.4, 6.0, 4.8, LIGHT_GRAY)
txt(s, "Channels Affected", 7.0, 1.55, 5.6, 0.5, size=18, bold=True, color=DARK)
channels = [
    ("Email Service",      "Retirement",      RED),
    ("SMS API",            "Retirement",      RED),
    ("Chat SDK",           "Retirement",      RED),
    ("Calling SDK",        "Breaking Change", ORANGE),
    ("Phone Numbers SDK",  "Retirement",      RED),
]
for i, (ch, status, col) in enumerate(channels):
    y = 2.2 + i * 0.75
    rect(s, 7.0, y, 0.15, 0.4, col)
    txt(s, ch,     7.25, y, 3.2, 0.4, size=13, bold=True, color=DARK)
    txt(s, status, 10.5, y, 2.0, 0.4, size=12, color=col)
txt(s, "Retirement date: March 31, 2029", 0.5, 6.5, 12, 0.5, size=12, color=GRAY_TEXT)

# ── Slide 3: The Solution ────────────────────────────────────────────────────
s = blank_slide(prs)
bg(s, WHITE)
rect(s, 0, 0, 13.33, 1.1, BLUE)
txt(s, "The Solution", 0.5, 0.15, 12, 0.8, size=32, bold=True, color=WHITE)
txt(s, "ACS Transition Agent automates the entire assessment workflow using AI agent skills and Azure CLI. No PowerShell required.", 0.5, 1.3, 12.3, 0.7, size=15, color=DARK)
steps = [
    ("1", "Authenticate",       "az login - verify Azure session"),
    ("2", "Select Subscription","Choose one or all subscriptions to scan"),
    ("3", "Discover Resources", "Find all ACS CommunicationServices resources"),
    ("4", "Detect Channels",    "Fast (Email/Phone) or Full (all 5 via Azure Monitor)"),
    ("5", "Analyze Impact",     "Map channels to migration guides and retirement dates"),
    ("6", "Generate Reports",   "CSV + Markdown + JSON saved to ./exports/"),
]
for i, (num, title, desc) in enumerate(steps):
    col = i % 3
    row = i // 3
    x = 0.4 + col * 4.3
    y = 2.3 + row * 2.2
    rect(s, x, y, 3.9, 1.85, LIGHT_GRAY)
    rect(s, x, y, 0.55, 0.55, BLUE)
    txt(s, num,   x,        y,        0.55, 0.55, size=18, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    txt(s, title, x + 0.65, y + 0.05, 3.1,  0.5,  size=13, bold=True, color=DARK)
    txt(s, desc,  x + 0.15, y + 0.65, 3.6,  1.0,  size=12, color=DARK)

# ── Slide 4: How to Use ──────────────────────────────────────────────────────
s = blank_slide(prs)
bg(s, WHITE)
rect(s, 0, 0, 13.33, 1.1, BLUE)
txt(s, "How to Use", 0.5, 0.15, 12, 0.8, size=32, bold=True, color=WHITE)
rect(s, 0.5, 1.3, 5.9, 5.5, LIGHT_GRAY)
txt(s, "Prerequisites", 0.7, 1.45, 5.5, 0.5, size=18, bold=True, color=DARK)
prereqs = [
    "Git",
    "Azure CLI  (learn.microsoft.com/cli/azure)",
    "GitHub Copilot or Claude Code",
    "Reader + Monitoring Reader on target subscription(s)",
]
for i, p in enumerate(prereqs):
    rect(s, 0.7, 2.1 + i * 0.7, 0.12, 0.4, BLUE)
    txt(s, p, 0.95, 2.05 + i * 0.7, 5.1, 0.5, size=13, color=DARK)
rect(s, 6.9, 1.3, 5.9, 5.5, LIGHT_GRAY)
txt(s, "Run the Scan", 7.1, 1.45, 5.5, 0.5, size=18, bold=True, color=DARK)
txt(s, "1.  Clone the repo\n\n2.  Open in VS Code with GitHub Copilot\n    or Claude Code\n\n3.  Type in chat:\n\n         Run an ACS deprecation scan\n\n4.  Follow the interactive prompts\n\n5.  Reports saved to  ./exports/", 7.1, 2.1, 5.5, 4.5, size=13, color=DARK)

# ── Slide 5: Migration Targets ───────────────────────────────────────────────
s = blank_slide(prs)
bg(s, WHITE)
rect(s, 0, 0, 13.33, 1.1, BLUE)
txt(s, "Migration Targets", 0.5, 0.15, 12, 0.8, size=32, bold=True, color=WHITE)
rows = [
    ("Email Service",      "Microsoft 365 High-Volume Email (HVE)",                           "aka.ms/acs-email-migration"),
    ("SMS API",            "Port numbers to a third-party SMS provider",                       "aka.ms/acs-sms-migration"),
    ("Chat SDK",           "Microsoft Teams Chat via Microsoft Graph APIs",                    "aka.ms/acs-chat-migration"),
    ("Calling SDK",        "Microsoft Teams (Phone Extensibility, Meeting Interop, Click-2-Call)", "aka.ms/acs-calling-migration"),
    ("Phone Numbers SDK",  "Port to Teams Phone Extensibility or third-party provider",        "aka.ms/acs-phone-migration"),
]
rect(s, 0.4, 1.2, 12.5, 0.5, BLUE)
txt(s, "Channel",                  0.5,  1.25, 3.0, 0.4, size=13, bold=True, color=WHITE)
txt(s, "Recommended Migration Path", 3.6, 1.25, 5.5, 0.4, size=13, bold=True, color=WHITE)
txt(s, "Guide",                    9.2,  1.25, 3.5, 0.4, size=13, bold=True, color=WHITE)
for i, (ch, path, guide) in enumerate(rows):
    y = 1.85 + i * 0.9
    row_color = LIGHT_GRAY if i % 2 == 0 else WHITE
    rect(s, 0.4, y, 12.5, 0.84, row_color)
    txt(s, ch,    0.5, y + 0.08, 3.0, 0.65, size=13, bold=True, color=DARK)
    txt(s, path,  3.6, y + 0.08, 5.4, 0.65, size=12, color=DARK)
    txt(s, guide, 9.2, y + 0.08, 3.5, 0.65, size=11, color=BLUE)
txt(s, "Full guide: aka.ms/acs-retirement-and-breaking-changes-guide", 0.4, 6.8, 12, 0.4, size=12, color=GRAY_TEXT)

# ── Slide 6: Get Started ─────────────────────────────────────────────────────
s = blank_slide(prs)
bg(s, DARK)
rect(s, 0, 0, 0.35, 7.5, BLUE)
txt(s, "Get Started", 0.6, 1.6, 12, 1.0, size=40, bold=True, color=WHITE)
txt(s, "github.com/jameelaesa/ACS-Transition-Agent-v0", 0.6, 2.8, 12, 0.6, size=20, color=LINK_BLUE)
txt(s, 'Type  "Run an ACS deprecation scan"  in any supported AI agent to begin.', 0.6, 3.7, 11, 0.6, size=18, color=RGBColor(200, 200, 200))
txt(s, "Retirement deadline: March 31, 2029", 0.6, 5.5, 10, 0.5, size=14, color=GRAY_TEXT)

out = "c:/Users/jameelaesa/ACS-Transition-Agent-v0/exports/ACS-Transition-Agent.pptx"
prs.save(out)
print(f"Saved: {out}")
